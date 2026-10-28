"""Convierte un reporte (ver datos.py) en archivo PDF o Excel, con el logo y el estilo del sistema.

Además de textos y números, una celda puede ser un diccionario:
    {"pill": "Suficiente", "tono": "ok"}          etiqueta de color (ok verde, wa amarillo, no rojo, ne gris, bl azul)
    {"barra": 73, "tono": "no", "texto": "73 %"}  barrita de porcentaje con su color
    {"t": "Leche cruda", "sub": "Materia prima"}  texto principal en negrita y una segunda línea gris
    {"t": "↑ 1.300", "tono": "ok"}               texto de color
Una sección puede ser una tabla o un gráfico de barras ({"tipo": "barras", "items": [(texto, valor)], "tono": ...}).
"""
from io import BytesIO

from django.conf import settings
from django.http import HttpResponse
from django.utils import timezone
from django.utils.text import slugify

SISTEMA = "Sistema de Control Estadístico y Trazabilidad en Procesos Lácteos"
AZUL = "16427A"
LOGO = settings.BASE_DIR / "static" / "img" / "logo.jpg"
FUENTES = settings.BASE_DIR / "static" / "fuentes"

# fondo, texto (como las etiquetas de la página)
TONOS = {
    "ok": ("DDF1E4", "165C37"),
    "wa": ("FDEBC8", "7A4A00"),
    "no": ("FBDCD7", "951B10"),
    "ne": ("E6E9EE", "3E4A57"),
    "bl": ("DCE7F6", "1D4F91"),
}
BARRAS = {"ok": "1F8A4C", "wa": "D98A00", "no": "C73A2E", "ne": "8C97A3", "bl": "3E7CC9",
          "azul": "16427A", "celeste": "3E7CC9"}
TEXTO_TONO = {"ok": "165C37", "no": "951B10", "wa": "7A4A00", "ne": "5B6662", "bl": "1D4F91"}


def _nombre_archivo(rep, ext):
    return f"{slugify(rep['titulo'])}_{timezone.localdate():%Y-%m-%d}.{ext}"


def _pie(usuario):
    nombre = usuario.get_full_name() or usuario.username
    return f"Generado por {nombre} ({usuario.get_rol_display()}) · Datos tomados del historial del sistema"


def _texto(v):
    """El texto plano de cualquier celda (para Excel y para medir)."""
    if isinstance(v, dict):
        if "pill" in v:
            return v["pill"]
        if "barra" in v:
            return v.get("texto", f"{v['barra']} %")
        if v.get("sub"):
            return f"{v['t']}\n{v['sub']}"
        return v.get("t", "")
    return "" if v is None else str(v)


def _es_grupo(fila):
    return isinstance(fila[0], str) and fila[0].startswith("» ")


# =====================================================================
# PDF
# =====================================================================
def _fuentes_pdf():
    """Usa DejaVu (tiene flechas ↑ ↓) si está en static/fuentes; si no, Helvetica."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    try:
        if "Prolac" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("Prolac", str(FUENTES / "DejaVuSans.ttf")))
            pdfmetrics.registerFont(TTFont("Prolac-B", str(FUENTES / "DejaVuSans-Bold.ttf")))
            from reportlab.lib.fonts import addMapping
            addMapping("Prolac", 0, 0, "Prolac")
            addMapping("Prolac", 1, 0, "Prolac-B")
            addMapping("Prolac", 0, 1, "Prolac")
            addMapping("Prolac", 1, 1, "Prolac-B")
        return "Prolac", "Prolac-B", True
    except Exception:
        return "Helvetica", "Helvetica-Bold", False


def a_pdf(rep, usuario):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.platypus import (CondPageBreak, Flowable, Image, KeepTogether, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    F, FB, unicode_ok = _fuentes_pdf()

    def limpio(texto):
        texto = str(texto)
        if not unicode_ok:
            texto = (texto.replace("↑", "+").replace("↓", "-").replace("→", "->").replace("·", "-")
                     .replace("…", "...").replace("≥", ">=").replace("≤", "<="))
        return texto

    def esc(texto):
        return limpio(texto).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    tablas = [s for s in rep["secciones"] if "cabeceras" in s]
    ancho_max = max([len(s["cabeceras"]) for s in tablas] or [1])
    pagina = landscape(letter) if ancho_max > 7 else letter
    margen = 14 * mm
    util = pagina[0] - 2 * margen - 12  # el marco de la página tiene 6 pt de relleno a cada lado

    c = lambda h: colors.HexColor("#" + h)  # noqa: E731
    azul, gris, tinta, borde = c(AZUL), c("3E5A7E"), c("14233A"), c("DCE4EE")

    def estilo(nombre, **kw):
        base = {"fontName": F, "fontSize": 7.8, "leading": 10, "textColor": tinta}
        base.update(kw)
        return ParagraphStyle(nombre, **base)

    est = {
        "titulo": estilo("t", fontName=FB, fontSize=20, leading=24, textColor=azul, alignment=TA_RIGHT),
        "sub": estilo("s", fontSize=8, leading=11, textColor=gris, alignment=TA_RIGHT),
        "sis": estilo("sis", fontName=FB, fontSize=8.5, leading=11, textColor=azul),
        "h2": estilo("h2", fontName=FB, fontSize=10, leading=13, textColor=azul, spaceBefore=12, spaceAfter=6),
        "celda": estilo("c"),
        "celda_d": estilo("cd", alignment=TA_RIGHT),
        "cab": estilo("cab", fontName=FB, fontSize=7.4, leading=9.2),
        "cab_d": estilo("cabd", fontName=FB, fontSize=7.4, leading=9.2, alignment=TA_RIGHT),
        "grupo": estilo("g", fontName=FB, textColor=azul),
        "grupo_d": estilo("gd", fontName=FB, textColor=azul, alignment=TA_RIGHT),
        "nota": estilo("n", fontSize=8, leading=11.5, textColor=gris),
        "barra_lab": estilo("bl", fontSize=8.5, leading=11),
    }

    class Etiqueta(Flowable):
        """Etiqueta redondeada de color (Pasó, Bajo el mínimo, 3 días...)."""

        def __init__(self, texto, tono="ne", borde_color=None, negrita=True, tam=7.2):
            super().__init__()
            self.texto, self.tam = limpio(texto), tam
            self.fondo, self.color = TONOS.get(tono, TONOS["ne"])
            self.fuente = FB if negrita else F
            self.borde_color = borde_color
            self.ancho = stringWidth(self.texto, self.fuente, tam) + 12
            self.alto = tam + 6

        def wrap(self, aw, ah):
            return min(self.ancho, aw), self.alto

        def draw(self):
            cv = self.canv
            cv.setFillColor(c(self.fondo))
            if self.borde_color:
                cv.setStrokeColor(c(self.borde_color))
                cv.setLineWidth(0.6)
            cv.roundRect(0, 0, self.ancho, self.alto, self.alto / 2, stroke=1 if self.borde_color else 0, fill=1)
            cv.setFillColor(c(self.color))
            cv.setFont(self.fuente, self.tam)
            cv.drawString(6, 3.2, self.texto)

    class Barra(Flowable):
        """Barrita de porcentaje con el número al lado."""

        def __init__(self, pct, tono, texto=None, ancho=None, alto=5.5, ancho_texto=30):
            super().__init__()
            self.pct = max(0, min(100, float(pct or 0)))
            self.color = BARRAS.get(tono, BARRAS["bl"])
            self.texto = limpio(texto) if texto is not None else None
            self.ancho_fijo, self.alto_barra, self.ancho_texto = ancho, alto, (ancho_texto if texto else 0)

        def wrap(self, aw, ah):
            self.ancho = self.ancho_fijo or aw
            return self.ancho, max(self.alto_barra, 9)

        def draw(self):
            cv = self.canv
            largo = self.ancho - self.ancho_texto - (4 if self.texto else 0)
            y = 1.8
            cv.setFillColor(c("E9EEF4"))
            cv.roundRect(0, y, largo, self.alto_barra, 1.5, stroke=0, fill=1)
            if self.pct > 0:
                cv.setFillColor(c(self.color))
                cv.roundRect(0, y, max(largo * self.pct / 100, 2), self.alto_barra, 1.5, stroke=0, fill=1)
            if self.texto:
                cv.setFillColor(tinta)
                cv.setFont(FB, 7.8)
                cv.drawRightString(self.ancho, y + 0.3, self.texto)

    class Kpi(Flowable):
        """Recuadro de número clave."""

        def __init__(self, k, ancho):
            super().__init__()
            self.k, self.ancho = k, ancho
            pequeno = k.get("pequeno")
            self.p_valor = Paragraph(esc(k["valor"]), estilo("kv", fontName=FB, fontSize=11 if pequeno else 17,
                                                           leading=14 if pequeno else 20))
            self.p_texto = Paragraph(esc(k["texto"]), estilo("kt", fontSize=7.6, leading=9.5, textColor=gris))
            cambio = k.get("cambio")
            self.p_cambio = None
            if cambio:
                color = TEXTO_TONO.get({"bueno": "ok", "malo": "no"}.get(cambio["tono"], "bl"), "3E5A7E")
                self.p_cambio = Paragraph(f"<font color='#{color}'><b>{esc(cambio['texto'])}</b></font>",
                                          estilo("kc", fontSize=7.2, leading=9))

        def wrap(self, aw, ah):
            interior = self.ancho - 20
            self.h1 = self.p_valor.wrap(interior, 100)[1]
            self.h2 = self.p_texto.wrap(interior, 100)[1]
            self.h3 = self.p_cambio.wrap(interior, 100)[1] if self.p_cambio else 0
            self.alto = max(18 + self.h1 + self.h2 + self.h3 + 2, getattr(self, "alto_min", 0))
            return self.ancho, self.alto

        def draw(self):
            cv = self.canv
            cv.setStrokeColor(borde)
            cv.setFillColor(colors.white)
            cv.setLineWidth(0.8)
            cv.roundRect(0, 0, self.ancho, self.alto, 6, stroke=1, fill=1)
            y = self.alto - 9 - self.h1
            self.p_valor.drawOn(cv, 10, y)
            y -= self.h2 + 1
            self.p_texto.drawOn(cv, 10, y)
            if self.p_cambio:
                y -= self.h3
                self.p_cambio.drawOn(cv, 10, y)

    def celda(v, derecha=False, ancho=None, grupo=False):
        if isinstance(v, dict):
            if "pill" in v:
                return Etiqueta(v["pill"], v.get("tono", "ne"))
            if "barra" in v:
                return Barra(v["barra"], v.get("tono", "bl"), v.get("texto", f"{v['barra']} %"), ancho=ancho)
            t = f"<b>{esc(v.get('t', ''))}</b>" if v.get("b", True) else esc(v.get("t", ""))
            if v.get("tono"):
                t = f"<font color='#{TEXTO_TONO.get(v['tono'], '14233A')}'>{t}</font>"
            if v.get("sub"):
                t += f"<br/><font color='#3E5A7E' size='7'>{esc(v['sub'])}</font>"
            return Paragraph(t, est["celda_d" if derecha else "celda"])
        if grupo:
            return Paragraph(esc(v), est["grupo_d" if derecha else "grupo"])
        return Paragraph(esc("" if v is None else v), est["celda_d" if derecha else "celda"])

    def _grafico(sec, titulo):
        """Gráfico de dona (partes de un todo) o de líneas (cómo cambia en el tiempo)."""
        from reportlab.graphics.charts.doughnut import Doughnut
        from reportlab.graphics.charts.linecharts import HorizontalLineChart
        from reportlab.graphics.shapes import Circle, Drawing, Line, Rect, String
        from reportlab.graphics.widgets.markers import makeMarker

        bloque = [titulo]
        if sec["tipo"] == "dona":
            partes = [p for p in sec["partes"] if p[1] > 0]
            total = sum(p[1] for p in sec["partes"])
            lado = 130
            d = Drawing(lado, lado)
            if partes:
                dona = Doughnut()
                dona.x = dona.y = 4
                dona.width = dona.height = lado - 8
                dona.data = [p[1] for p in partes]
                dona.innerRadiusFraction = 0.6
                dona.strokeWidth = 1.5
                dona.strokeColor = colors.white
                dona.slices.strokeColor = colors.white
                dona.slices.strokeWidth = 1.5
                dona.startAngle = 90
                dona.direction = "clockwise"
                for i, p in enumerate(partes):
                    dona.slices[i].fillColor = c(BARRAS.get(p[2], BARRAS["ne"]))
                d.add(dona)
            else:
                d.add(Circle(lado / 2, lado / 2, lado / 2 - 4, fillColor=c("E9EEF4"), strokeColor=None))
                d.add(Circle(lado / 2, lado / 2, (lado / 2 - 4) * 0.6, fillColor=colors.white, strokeColor=None))
            d.add(String(lado / 2, lado / 2 + 2, limpio(sec.get("centro", str(total))), fontName=FB, fontSize=20,
                         fillColor=tinta, textAnchor="middle"))
            d.add(String(lado / 2, lado / 2 - 12, limpio(sec.get("centro_texto", "")), fontName=F, fontSize=7.5,
                         fillColor=gris, textAnchor="middle"))
            filas_l = []
            for nombre, valor, tono, *extra in sec["partes"]:
                punto = Drawing(10, 10)
                punto.add(Circle(5, 5, 4.5, fillColor=c(BARRAS.get(tono, BARRAS["ne"])), strokeColor=None))
                pct = f"{round(100 * valor / total)} %" if total else "0 %"
                filas_l.append([punto, Paragraph(f"<b>{esc(nombre)}</b>" + (f"<br/><font color='#3E5A7E' size='7'>{esc(extra[0])}</font>" if extra else ""),
                                                  estilo("dl", fontSize=8.5, leading=11)),
                                Paragraph(f"<b>{esc(valor)}</b>", estilo("dv", fontName=FB, fontSize=11, leading=13, alignment=TA_RIGHT)),
                                Paragraph(pct, estilo("dp", fontSize=8, leading=11, textColor=gris, alignment=TA_RIGHT))])
            ley = Table(filas_l, colWidths=[16, util - lado - 150, 40, 44])
            ley.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 0), (-1, -2), 0.4, c("E9EEF4")),
                                     ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                                     ("LEFTPADDING", (0, 0), (-1, -1), 2)]))
            t = Table([[d, "", ley]], colWidths=[lado, 30, util - lado - 30])
            t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                   ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
            bloque.append(t)
        else:
            etiquetas = sec["etiquetas"]
            series = sec["series"]  # [(nombre, tono, [valores])]
            alto, ancho = 150, util
            d = Drawing(ancho, alto)
            g = HorizontalLineChart()
            g.x, g.y, g.width, g.height = 28, 26, ancho - 40, alto - 44
            g.data = [s_[2] for s_ in series]
            g.joinedLines = 1
            g.categoryAxis.categoryNames = [limpio(e) for e in etiquetas]
            g.categoryAxis.labels.fontName = F
            g.categoryAxis.labels.fontSize = 7.5
            g.categoryAxis.labels.fillColor = gris
            g.categoryAxis.strokeColor = c("C9D5E4")
            g.categoryAxis.tickDown = 0
            g.valueAxis.valueMin = 0
            tope = max([v for s_ in series for v in s_[2]] or [1])
            paso = max(1, round(tope / 4 / 5) * 5) if tope > 10 else max(1, round(tope / 4) or 1)
            g.valueAxis.valueMax = paso * (int(tope / paso) + 1)
            g.valueAxis.valueStep = paso
            g.valueAxis.labels.fontName = F
            g.valueAxis.labels.fontSize = 7
            g.valueAxis.labels.fillColor = gris
            g.valueAxis.strokeColor = None
            g.valueAxis.visibleGrid = 1
            g.valueAxis.gridStrokeColor = c("E9EEF4")
            g.valueAxis.gridStrokeWidth = 0.6
            for i, (nombre, tono, valores) in enumerate(series):
                col = c(BARRAS.get(tono, BARRAS["bl"]))
                g.lines[i].strokeColor = col
                g.lines[i].strokeWidth = 2.2
                g.lines[i].symbol = makeMarker("FilledCircle", size=5, fillColor=col, strokeColor=colors.white)
            g.lineLabelFormat = "%d"
            g.lineLabels.fontName = FB
            g.lineLabels.fontSize = 7.5
            g.lineLabels.dy = 6
            g.lineLabels.fillColor = tinta
            if len(series) == 2:  # la etiqueta del valor menor va debajo del punto para que no se monten
                for j in range(len(etiquetas)):
                    a_, b_ = series[0][2][j], series[1][2][j]
                    abajo = 1 if a_ >= b_ else 0
                    g.lineLabels[(abajo, j)].dy = -13
            d.add(g)
            bloque.append(d)
            partes, anchos_l = [], []
            for nombre, tono, _v in series:
                linea = Drawing(22, 8)
                linea.add(Line(0, 4, 16, 4, strokeColor=c(BARRAS.get(tono, BARRAS["bl"])), strokeWidth=2.2))
                linea.add(Circle(8, 4, 2.6, fillColor=c(BARRAS.get(tono, BARRAS["bl"])), strokeColor=colors.white))
                partes += [linea, Paragraph(esc(nombre), estilo("lg", fontSize=8, leading=10))]
                anchos_l += [22, stringWidth(limpio(nombre), F, 8) + 18]
            ley = Table([partes], colWidths=anchos_l, hAlign="LEFT")
            ley.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
            bloque += [Spacer(1, 2), ley]
        if sec.get("nota"):
            bloque += [Spacer(1, 5), Paragraph(esc(sec["nota"]), est["nota"])]
        return bloque

    historia = []

    # ---- Encabezado: logo + sistema | título -------------------------
    try:
        logo = Image(str(LOGO), width=19 * mm, height=19 * mm, kind="proportional")
    except Exception:
        logo = Spacer(1, 1)
    izq = Table([[logo, Paragraph(f"{esc(SISTEMA)}<br/><font name='{F}' color='#3E5A7E' size='8'>Inventario de lácteos</font>",
                                  est["sis"])]], colWidths=[22 * mm, 68 * mm])
    izq.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    der = [Paragraph(esc(rep["titulo"]), est["titulo"]),
           Paragraph(f"{esc(rep['subtitulo'])}<br/>Generado el {timezone.localtime():%d/%m/%Y a las %H:%M}", est["sub"])]
    cab = Table([[izq, der]], colWidths=[util * 0.5, util * 0.5])
    cab.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 0), (-1, 0), 1.6, azul),
                             ("BOTTOMPADDING", (0, 0), (-1, -1), 9), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                             ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    historia += [cab, Spacer(1, 10)]

    # ---- Filtros como fichas -----------------------------------------
    if rep.get("filtros"):
        fichas, anchos = [], []
        for t in rep["filtros"]:
            if ":" in t:
                a, b = t.split(":", 1)
                html = f"<b>{esc(a)}:</b>{esc(b)}"
            else:
                html = esc(t)
            p = Paragraph(html, estilo("fi", fontSize=7.6, leading=9.5))
            w = stringWidth(limpio(t), FB, 7.6) + 14
            fichas.append(p)
            anchos.append(w)
        t = Table([fichas], colWidths=anchos, hAlign="LEFT")
        estilo_f = [("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]
        t.setStyle(TableStyle(estilo_f))

        class Fichas(Flowable):
            def __init__(self, tabla, anchos):
                super().__init__()
                self.tabla, self.anchos = tabla, anchos

            def wrap(self, aw, ah):
                self.w, self.h = self.tabla.wrap(aw, ah)
                return self.w, self.h

            def draw(self):
                cv, x = self.canv, 0
                for w in self.anchos:
                    cv.setFillColor(c("F2F6FB"))
                    cv.setStrokeColor(c("C9D5E4"))
                    cv.setLineWidth(0.6)
                    cv.roundRect(x + 1, 0.5, w - 6, self.h - 1, 4, stroke=1, fill=1)
                    x += w
                self.tabla.drawOn(cv, 0, 0)
        # cada ficha deja un espacio de 4 pt a la derecha
        historia += [Fichas(t, anchos), Spacer(1, 10)]

    # ---- Números clave -----------------------------------------------
    if rep.get("kpis"):
        n = len(rep["kpis"])
        hueco = 8
        ancho_k = (util - hueco * (n - 1)) / n
        cajas = [Kpi(k, ancho_k) for k in rep["kpis"]]
        fila = []
        for i, caja in enumerate(cajas):
            fila.append(caja)
            if i < n - 1:
                fila.append("")
        anchos = []
        for i in range(n):
            anchos.append(ancho_k)
            if i < n - 1:
                anchos.append(hueco)
        t = Table([fila], colWidths=anchos)
        t.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                               ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                               ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        altura = max(caja.wrap(ancho_k, 200)[1] for caja in cajas)
        for caja in cajas:
            caja.alto_min = altura  # todas las cajas con la misma altura
        historia += [t, Spacer(1, 4)]

    # ---- Secciones ---------------------------------------------------
    for sec in rep["secciones"]:
        titulo = Paragraph(esc(sec["titulo"].upper()), est["h2"])
        if sec.get("tipo") in ("dona", "lineas"):
            historia.append(KeepTogether(_grafico(sec, titulo)))
            continue
        if sec.get("tipo") == "barras":
            items = sec["items"]
            series = sec.get("series")  # [(nombre, tono), ...] para barras dobles
            valores = [v for it in items for v in (it[1] if series else [it[1]])]
            tope = sec.get("tope") or max(valores or [1]) or 1
            est_v = estilo("bv", fontSize=8.5, leading=11, alignment=TA_RIGHT)
            filas = []
            for it in items:
                if series:
                    textos = it[2] if len(it) > 2 else it[1]
                    barras_it = Table([[Barra(100 * v / tope, tono, None, alto=5.5),
                                        Paragraph(f"<b>{esc(tx)}</b>", est_v)]
                                       for v, tx, (_, tono) in zip(it[1], textos, series)],
                                      colWidths=[util * 0.6, util * 0.1])
                    barras_it.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                                   ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                                                   ("TOPPADDING", (0, 0), (-1, -1), 0.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 0.5)]))
                    filas.append([Paragraph(esc(it[0]), est["barra_lab"]), barras_it])
                else:
                    tono = it[3] if len(it) > 3 else sec.get("tono", "no")
                    filas.append([Paragraph(esc(it[0]), est["barra_lab"]),
                                  Barra(100 * min(it[1], tope) / tope, tono, None, alto=6.5),
                                  Paragraph(f"<b>{esc(it[2] if len(it) > 2 else it[1])}</b>", est_v)])
            if not filas:
                filas = [[Paragraph(sec.get("vacio", "Sin datos en el periodo."), est["nota"]), ""] + ([] if series else [""])]
            anchos_b = [util * 0.3, util * 0.7] if series else [util * 0.3, util * 0.55, util * 0.15]
            t = Table(filas, colWidths=anchos_b)
            t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                   ("TOPPADDING", (0, 0), (-1, -1), 2 if series else 1.5),
                                   ("BOTTOMPADDING", (0, 0), (-1, -1), 2 if series else 1.5)]
                                  + ([("LINEBELOW", (0, 0), (-1, -2), 0.4, c("E9EEF4"))] if series else [])))
            bloque_ley = []
            ley_items = sec.get("leyenda") or ([(n, tono, "") for n, tono in series] if series else [])
            if ley_items:
                partes, anchos_l = [], []
                for texto, tono, explica in ley_items:
                    partes += [Etiqueta(texto, tono), Paragraph(esc(explica), estilo("ly", fontSize=7.6, leading=9.5, textColor=gris))]
                    anchos_l += [stringWidth(limpio(texto), FB, 7.2) + 16, stringWidth(limpio(explica), F, 7.6) + 14]
                ley = Table([partes], colWidths=anchos_l, hAlign="LEFT")
                ley.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                         ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
                bloque_ley = [Spacer(1, 5), ley]
            bloque = [titulo, t] + bloque_ley
            if sec.get("nota"):
                bloque += [Spacer(1, 4), Paragraph(esc(sec["nota"]), est["nota"])]
            historia.append(KeepTogether(bloque))
            continue

        cabeceras = sec["cabeceras"]
        alin = sec.get("alinear") or "l" * len(cabeceras)
        if sec.get("ancho"):
            total = sum(sec["ancho"])
            anchos = [util * a / total for a in sec["ancho"]]
        else:
            anchos = [util / len(cabeceras)] * len(cabeceras)
        datos = [[Paragraph(esc(h), est["cab_d" if alin[j] == "r" else "cab"]) for j, h in enumerate(cabeceras)]]
        grupos = []
        for i, fila in enumerate(sec["filas"], 1):
            grupo = _es_grupo(fila)
            if grupo:
                grupos.append(i)
                fila = [fila[0][2:]] + list(fila[1:])
            datos.append([celda(v, alin[j] == "r", anchos[j] - 12, grupo) for j, v in enumerate(fila)])
        if not sec["filas"]:
            datos.append([Paragraph(sec.get("vacio", "Sin datos en el periodo."), est["nota"])] + [""] * (len(cabeceras) - 1))
        t = Table(datos, colWidths=anchos, repeatRows=1)
        estilo_t = [("BACKGROUND", (0, 0), (-1, 0), c("EDF2F8")),
                    ("LINEBELOW", (0, 0), (-1, 0), 0.8, c("C9D5E4")),
                    ("LINEBELOW", (0, 1), (-1, -1), 0.4, c("E9EEF4")),
                    ("VALIGN", (0, 0), (-1, 0), "BOTTOM"),
                    ("VALIGN", (0, 1), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 4.5),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6)]
        if not sec["filas"]:
            estilo_t.append(("SPAN", (0, 1), (-1, 1)))
        for g in grupos:
            estilo_t.append(("BACKGROUND", (0, g), (-1, g), c("F5F8FC")))
            fila_g = sec["filas"][g - 1]
            hasta = 0
            while hasta + 1 < len(fila_g) and fila_g[hasta + 1] in ("", None):
                hasta += 1
            if hasta:
                estilo_t.append(("SPAN", (0, g), (hasta, g)))
        t.setStyle(TableStyle(estilo_t))
        historia += [CondPageBreak(40 * mm), titulo, t]
        if sec.get("leyenda"):
            filas_l, fila_l, anchos_fila, usado = [], [], [], 0
            for texto, tono, explica in sec["leyenda"]:
                w1 = stringWidth(limpio(texto), FB, 7.2) + 16
                w2 = stringWidth(limpio(explica), F, 7.6) + 14
                if fila_l and usado + w1 + w2 > util:
                    filas_l.append((fila_l, anchos_fila))
                    fila_l, anchos_fila, usado = [], [], 0
                fila_l += [Etiqueta(texto, tono), Paragraph(esc(explica), estilo("ly", fontSize=7.6, leading=9.5, textColor=gris))]
                anchos_fila += [w1, w2]
                usado += w1 + w2
            if fila_l:
                filas_l.append((fila_l, anchos_fila))
            ley = []
            for fl, an in filas_l:
                tl = Table([fl], colWidths=an, hAlign="LEFT")
                tl.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                        ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 1.5),
                                        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5)]))
                ley.append(tl)
            historia += [Spacer(1, 7)] + ley
        if sec.get("nota"):
            historia += [Spacer(1, 5), Paragraph(esc(sec["nota"]), est["nota"])]

    if rep.get("nota_final"):
        texto = esc(rep["nota_final"])
        if ":" in texto[:40]:
            a, b = texto.split(":", 1)
            texto = f"<font color='#{AZUL}'><b>{a}:</b></font>{b}"
        historia += [Spacer(1, 12), Paragraph(texto, est["nota"])]

    pie = limpio(_pie(usuario))

    class CanvasNumerado(rl_canvas.Canvas):
        """Escribe "Página X de Y" cuando ya sabe cuántas páginas hay."""

        def __init__(self, *a, **kw):
            super().__init__(*a, **kw)
            self._paginas = []

        def showPage(self):
            self._paginas.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._paginas)
            for estado in self._paginas:
                self.__dict__.update(estado)
                self.saveState()
                self.setStrokeColor(borde)
                self.setLineWidth(0.6)
                self.line(margen + 6, 12 * mm, pagina[0] - margen - 6, 12 * mm)
                self.setFont(F, 7)
                self.setFillColor(gris)
                self.drawString(margen + 6, 8 * mm, pie)
                self.drawRightString(pagina[0] - margen - 6, 8 * mm, f"Página {self._pageNumber} de {total}")
                self.restoreState()
                super().showPage()
            super().save()

    salida = BytesIO()
    doc = SimpleDocTemplate(salida, pagesize=pagina, leftMargin=margen, rightMargin=margen,
                            topMargin=14 * mm, bottomMargin=18 * mm, title=rep["titulo"], author=SISTEMA)
    doc.build(historia, canvasmaker=CanvasNumerado)
    respuesta = HttpResponse(salida.getvalue(), content_type="application/pdf")
    respuesta["Content-Disposition"] = f'attachment; filename="{_nombre_archivo(rep, "pdf")}"'
    return respuesta


# =====================================================================
# Excel
# =====================================================================
def a_excel(rep, usuario):
    from openpyxl import Workbook
    from openpyxl.chart import DoughnutChart, LineChart, Reference
    from openpyxl.chart.series import DataPoint
    from openpyxl.formatting.rule import DataBarRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    libro = Workbook()
    hoja = libro.active
    hoja.title = rep["titulo"][:31]
    hoja.sheet_view.showGridLines = False
    borde = Border(bottom=Side(style="thin", color="DCE4EE"))
    caja = Border(*(Side(style="thin", color="DCE4EE"),) * 4)

    try:
        from openpyxl.drawing.image import Image
        img = Image(str(LOGO))
        img.width, img.height = 64, 64
        hoja.add_image(img, "A1")
        hoja.row_dimensions[1].height = 26
        hoja.row_dimensions[2].height = 26
        col_texto = "B"
    except Exception:
        col_texto = "A"

    hoja[f"{col_texto}1"] = rep["titulo"]
    hoja[f"{col_texto}1"].font = Font(size=18, bold=True, color=AZUL)
    hoja[f"{col_texto}2"] = f"{rep['subtitulo']} · {SISTEMA}"
    hoja[f"{col_texto}2"].font = Font(size=10, color="3E5A7E")
    fila = 4
    for j, t in enumerate(rep.get("filtros", []), 1):
        celda = hoja.cell(fila, j, t)
        celda.font = Font(size=9, bold=True, color="14233A")
        celda.fill = PatternFill("solid", fgColor="F2F6FB")
        celda.border = caja
    fila += 2

    anchos = {}
    for i, k in enumerate(rep.get("kpis", [])):
        col = 1 + i * 2
        c = hoja.cell(fila, col, k["valor"])
        c.font = Font(size=16, bold=True, color="14233A")
        hoja.cell(fila + 1, col, k["texto"]).font = Font(size=9, color="3E5A7E")
        if k.get("cambio"):
            color = {"bueno": "165C37", "malo": "951B10"}.get(k["cambio"]["tono"], "3E5A7E")
            hoja.cell(fila + 2, col, k["cambio"]["texto"]).font = Font(size=8, bold=True, color=color)
        anchos[col] = max(anchos.get(col, 0), 18)
    fila += 4

    def poner(fila_n, col, v, grupo=False):
        c = hoja.cell(fila_n, col)
        c.border = borde
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if isinstance(v, dict) and "pill" in v:
            fondo, texto = TONOS.get(v.get("tono", "ne"), TONOS["ne"])
            c.value = v["pill"]
            c.fill = PatternFill("solid", fgColor=fondo)
            c.font = Font(bold=True, color=texto, size=10)
            c.alignment = Alignment(horizontal="center", vertical="top")
        elif isinstance(v, dict) and "barra" in v:
            c.value = float(v["barra"]) / 100
            c.number_format = "0%"
            c.font = Font(bold=True, color="14233A")
            color = BARRAS.get(v.get("tono", "bl"), BARRAS["bl"])
            hoja.conditional_formatting.add(
                c.coordinate, DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=color))
        elif isinstance(v, dict):
            c.value = _texto(v)
            c.font = Font(bold=v.get("b", True), color=TEXTO_TONO.get(v.get("tono"), "14233A"))
        else:
            c.value = v
            if grupo:
                c.font = Font(bold=True, color=AZUL)
        if grupo:
            c.fill = PatternFill("solid", fgColor="F5F8FC")
        anchos[col] = max(anchos.get(col, 0), min(max(len(x) for x in (_texto(v) or " ").split("\n")), 60))

    for sec in rep["secciones"]:
        hoja.cell(fila, 1, sec["titulo"].upper()).font = Font(size=11, bold=True, color=AZUL)
        fila += 1
        if sec.get("tipo") == "dona":
            inicio = fila
            for nombre, valor, tono, *extra in sec["partes"]:
                hoja.cell(fila, 1, nombre).font = Font(bold=True, color=TEXTO_TONO.get(tono, "14233A"))
                hoja.cell(fila, 2, valor).font = Font(bold=True)
                anchos[1] = max(anchos.get(1, 0), len(nombre))
                fila += 1
            grafico = DoughnutChart(holeSize=55)
            grafico.add_data(Reference(hoja, min_col=2, min_row=inicio, max_row=fila - 1))
            grafico.set_categories(Reference(hoja, min_col=1, min_row=inicio, max_row=fila - 1))
            serie = grafico.series[0]
            for i, (_n, _v, tono, *_e) in enumerate(sec["partes"]):
                punto = DataPoint(idx=i)
                punto.graphicalProperties.solidFill = BARRAS.get(tono, BARRAS["ne"])
                punto.graphicalProperties.line.solidFill = "FFFFFF"
                serie.dPt.append(punto)
            grafico.title = sec["titulo"]
            grafico.height, grafico.width = 7, 11
            hoja.add_chart(grafico, f"D{inicio}")
            fila = max(fila, inicio + 15)
            if sec.get("nota"):
                hoja.cell(fila, 1, sec["nota"]).font = Font(size=9, italic=True, color="3E5A7E")
                fila += 1
            fila += 1
            continue
        if sec.get("tipo") == "lineas":
            hoja.cell(fila, 1, "Periodo").font = Font(bold=True, color="FFFFFF")
            hoja.cell(fila, 1).fill = PatternFill("solid", fgColor=AZUL)
            for j, (nombre, tono, _v) in enumerate(sec["series"], 2):
                c = hoja.cell(fila, j, nombre)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor=AZUL)
                anchos[j] = max(anchos.get(j, 0), len(nombre))
            cab = fila
            fila += 1
            for k, etiqueta in enumerate(sec["etiquetas"]):
                hoja.cell(fila, 1, etiqueta)
                anchos[1] = max(anchos.get(1, 0), len(etiqueta))
                for j, (_n, _t, valores) in enumerate(sec["series"], 2):
                    hoja.cell(fila, j, valores[k])
                fila += 1
            if sec["etiquetas"]:
                grafico = LineChart()
                grafico.add_data(Reference(hoja, min_col=2, max_col=1 + len(sec["series"]), min_row=cab, max_row=fila - 1),
                                 titles_from_data=True)
                grafico.set_categories(Reference(hoja, min_col=1, min_row=cab + 1, max_row=fila - 1))
                for serie, (_n, tono, _v) in zip(grafico.series, sec["series"]):
                    serie.graphicalProperties.line.solidFill = BARRAS.get(tono, BARRAS["bl"])
                    serie.graphicalProperties.line.width = 28000
                    serie.marker.symbol = "circle"
                    serie.marker.graphicalProperties.solidFill = BARRAS.get(tono, BARRAS["bl"])
                    serie.smooth = False
                grafico.title = sec["titulo"]
                grafico.height, grafico.width = 7.5, 16
                hoja.add_chart(grafico, f"E{cab}")
                fila = max(fila, cab + 16)
            if sec.get("nota"):
                hoja.cell(fila, 1, sec["nota"]).font = Font(size=9, italic=True, color="3E5A7E")
                fila += 1
            fila += 1
            continue
        if sec.get("tipo") == "barras":
            series = sec.get("series")
            if series:
                for j, (nombre, tono) in enumerate(series, 2):
                    c = hoja.cell(fila, j, nombre)
                    c.font = Font(bold=True, color="FFFFFF")
                    c.fill = PatternFill("solid", fgColor=AZUL)
                fila += 1
            inicio = fila
            for it in sec["items"]:
                hoja.cell(fila, 1, it[0])
                anchos[1] = max(anchos.get(1, 0), len(it[0]))
                vals = it[1] if series else [it[1]]
                for j, v in enumerate(vals, 2):
                    celda = hoja.cell(fila, j, float(v))
                    celda.font = Font(bold=True)
                    if not series and len(it) > 3:  # color propio de cada barra
                        hoja.conditional_formatting.add(celda.coordinate, DataBarRule(
                            start_type="num", start_value=0, end_type="num", end_value=float(sec.get("tope") or max(x[1] for x in sec["items"]) or 1),
                            color=BARRAS.get(it[3], BARRAS["no"])))
                    anchos[j] = max(anchos.get(j, 0), 24)
                fila += 1
            if sec["items"]:
                tonos = [t for _, t in series] if series else [None]
                for j, tono in enumerate(tonos, 2):
                    if not series and len(sec["items"][0]) > 3:
                        continue
                    letra = get_column_letter(j)
                    color = BARRAS.get(tono or sec.get("tono", "no"), BARRAS["no"])
                    hoja.conditional_formatting.add(f"{letra}{inicio}:{letra}{fila - 1}", DataBarRule(
                        start_type="num", start_value=0, end_type="max", color=color))
            fila += 1
            continue
        for j, cab in enumerate(sec["cabeceras"], 1):
            c = hoja.cell(fila, j, cab)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=AZUL)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            anchos[j] = max(anchos.get(j, 0), len(str(cab)))
        fila += 1
        if not sec["filas"]:
            hoja.cell(fila, 1, sec.get("vacio", "Sin datos en el periodo.")).font = Font(italic=True, color="5B6662")
            fila += 1
        for f in sec["filas"]:
            grupo = _es_grupo(f)
            if grupo:
                f = [f[0][2:]] + list(f[1:])
            for j, valor in enumerate(f, 1):
                poner(fila, j, valor, grupo)
            fila += 1
        if sec.get("leyenda"):
            j = 1
            for texto, tono, explica in sec["leyenda"]:
                poner(fila, j, {"pill": texto, "tono": tono})
                hoja.cell(fila, j + 1, explica).font = Font(size=9, color="3E5A7E")
                j += 2
            fila += 1
        if sec.get("nota"):
            hoja.cell(fila, 1, sec["nota"]).font = Font(size=9, italic=True, color="3E5A7E")
            fila += 1
        fila += 1

    if rep.get("nota_final"):
        hoja.cell(fila, 1, rep["nota_final"]).font = Font(bold=True, color="14233A")
        fila += 2
    hoja.cell(fila, 1, f"{_pie(usuario)} · {timezone.localtime():%d/%m/%Y %H:%M}").font = Font(size=8, color="5B6662")

    for j, ancho in anchos.items():
        hoja.column_dimensions[get_column_letter(j)].width = max(11, ancho + 2)

    salida = BytesIO()
    libro.save(salida)
    respuesta = HttpResponse(salida.getvalue(),
                             content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    respuesta["Content-Disposition"] = f'attachment; filename="{_nombre_archivo(rep, "xlsx")}"'
    return respuesta
